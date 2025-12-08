
import os
import json
from datetime import date

from docx import Document

from utils_docx import (
    replace_placeholders,
    inserir_foto_por_placeholder,
    format_date_br,
    date_por_extenso,
)


def generate_ultrassom_report(dados: dict, template_path: str, output_dir: str) -> str:
    """Gera o relatório de Ultrassom a partir do template e dos dados fornecidos.

    Retorna o caminho completo do arquivo gerado.
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template não encontrado: {template_path}")

    # Garante diretório de saída
    os.makedirs(output_dir, exist_ok=True)

    numrel = str(dados.get("NUMRELATORIO", "")).strip()
    if not numrel:
        raise ValueError("O campo NUMRELATORIO é obrigatório.")

    # Trata datas
    data_insp = dados.get("DATA_INSP")
    if isinstance(data_insp, date):
        data_insp_br = format_date_br(data_insp)
        data_extenso = date_por_extenso(data_insp)
    else:
        data_insp_br = str(data_insp) if data_insp else ""
        data_extenso = ""

    dados_limpos = dict(dados)
    dados_limpos["DATA_INSP"] = data_insp_br
    dados_limpos["DATA_INSP_EXTENSO"] = data_extenso

    doc = Document(template_path)

    mapping: dict[str, str] = {}
    for k, v in dados_limpos.items():
        if k.startswith("FOTO_"):
            continue
        placeholder = "{{" + k + "}}"
        mapping[placeholder] = "" if v is None else str(v)

    # ---------------- CONCLUSÃO AUTOMÁTICA ----------------
    laudo = dados.get("LAUDO")
    laudo_ext = dados.get("LAUDO_EXTENSO") or ""

    if laudo == "A":
        conclusao = (
            f"CONCLUSÃO: Não Foram evidenciadas indicações relevantes, "
            f"estando o ensaio {laudo_ext} segundo o critério de aceitação da norma acima"
        )
    else:
        conclusao = (
            f"CONCLUSÃO: Foram evidenciadas indicações relevantes, "
            f"estando o ensaio {laudo_ext} segundo o critério de aceitação da norma acima"
        )

    mapping["{{CONCLUSAO}}"] = conclusao

    replace_placeholders(doc, mapping)

    # Inserção das fotos com tamanhos adequados
    foto1 = dados.get("FOTO_1")
    if foto1:
        inserir_foto_por_placeholder(doc, "{{FOTO_1}}", foto1, height_cm=10.0)

    foto2 = dados.get("FOTO_2")
    if foto2:
        inserir_foto_por_placeholder(doc, "{{FOTO_2}}", foto2, height_cm=4.0)

    foto3 = dados.get("FOTO_3")
    if foto3:
        inserir_foto_por_placeholder(doc, "{{FOTO_3}}", foto3, height_cm=4.0)

    file_name = f"Relatório {numrel}-US.docx"
    output_path = os.path.join(output_dir, file_name)
    doc.save(output_path)

    return output_path