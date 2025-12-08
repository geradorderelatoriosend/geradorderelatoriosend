import os
from datetime import date

from docx import Document

from utils_docx import (
    replace_placeholders,
    inserir_foto_por_placeholder,
    format_date_br,
    date_por_extenso,
)


def generate_capa_report(dados: dict, template_path: str, output_dir: str) -> str:
    """
    Gera a CAPA a partir do CAPA_TEMPLATE.docx.

    Espera, no mínimo, campos como:
      - NUMRELATORIO
      - EMPRESA, ENDEREÇO, BAIRRO, CIDADE, ESTADO, CEP
      - CONTATO, DDD, FONE, EMAIL
      - PECA_INSP, NUM_DESENHO, QUANTIDADE
      - DATA_INSP (date)  -> DATA_INSP e DATA_INSP_EXTENSO
      - TIPO_ENSAIO
      - FOTO_1 (foto da capa)

    Retorna o caminho absoluto do arquivo gerado.
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template de capa não encontrado: {template_path}")

    os.makedirs(output_dir, exist_ok=True)

    doc = Document(template_path)

    mapping: dict[str, str] = {}

    # Mapeia todos os campos de texto (exceto FOTO_*)
    for key, value in dados.items():
        if key.startswith("FOTO_"):
            continue

        ph = f"{{{{{key}}}}}"  # ex: "EMPRESA" -> "{{EMPRESA}}"

        if key == "DATA_INSP" and isinstance(value, date):
            mapping[ph] = format_date_br(value)
        elif key == "DATA_INSP_EXTENSO" and isinstance(value, date):
            # Se por acaso mandarem a data em si aqui
            mapping[ph] = date_por_extenso(value)
        else:
            mapping[ph] = "" if value is None else str(value)

    # Se tiver DATA_INSP e ainda não tiver DATA_INSP_EXTENSO preenchido,
    # gera automaticamente
    data_insp = dados.get("DATA_INSP")
    if isinstance(data_insp, date) and "{{DATA_INSP_EXTENSO}}" not in mapping:
        mapping["{{DATA_INSP_EXTENSO}}"] = date_por_extenso(data_insp)

    # Aplica todos os placeholders de texto
    replace_placeholders(doc, mapping)

    # FOTO_1 = foto da capa (altura maior)
    foto1 = dados.get("FOTO_1")
    if foto1:
        inserir_foto_por_placeholder(doc, "{{FOTO_1}}", foto1, height_cm=10.0)

    num = (dados.get("NUMRELATORIO") or "").strip() or "0000"
    filename = f"Relatório {num}-CAPA.docx"
    out_path = os.path.join(output_dir, filename)

    doc.save(out_path)
    return os.path.abspath(out_path)
