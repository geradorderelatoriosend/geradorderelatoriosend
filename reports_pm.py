import os
from datetime import date

from docx import Document

from utils_docx import (
    replace_placeholders,
    inserir_foto_por_placeholder,
    format_date_br,
    date_por_extenso,
)


def generate_pm_report(dados: dict, template_path: str, output_dir: str) -> str:
    """
    Gera o relatório de Partículas Magnéticas (PM) a partir do PM_TEMPLATE.docx.

    Campos usados no template:
      - NUMRELATORIO
      - EMPRESA, ENDEREÇO, BAIRRO, CIDADE, ESTADO, CEP, CONTATO, DDD, FONE, EMAIL
      - DATA_INSP, DATA_INSP_EXTENSO
      - PECA_INSP, NUM_DESENHO, QUANTIDADE
      - FAB_PARTICULA, VAL_PARTICULA, LOTE_PARTICULA
      - TEMPERATURA, COND_SUPERFICIAL
      - FOTO_1, FOTO_2
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template não encontrado: {template_path}")

    os.makedirs(output_dir, exist_ok=True)

    doc = Document(template_path)

    # ---------------- TEXTO (sem fotos) ----------------
    mapping: dict[str, str] = {}

    for key, value in dados.items():
        # Fotos são tratadas separadamente
        if key in ("FOTO_1", "FOTO_2"):
            continue

        ph = f"{{{{{key}}}}}"  # ex: "DATA_INSP" -> "{{DATA_INSP}}"

        if key == "DATA_INSP" and isinstance(value, date):
            mapping[ph] = format_date_br(value)
        else:
            mapping[ph] = "" if value is None else str(value)

    # Data por extenso
    data_insp = dados.get("DATA_INSP")
    if isinstance(data_insp, date):
        mapping["{{DATA_INSP_EXTENSO}}"] = date_por_extenso(data_insp)
    else:
        mapping["{{DATA_INSP_EXTENSO}}"] = ""

    # ---------------- CONCLUSÃO AUTOMÁTICA ----------------
    laudo = dados.get("LAUDO")
    laudo_ext = dados.get("LAUDO_EXTENSO") or ""

    if laudo == "A":
        conclusao = (
            f"CONCLUSÃO: Não Foram evidenciadas indicações relevantes, "
            f"estando o ensaio {laudo_ext} segundo o critério de aceitação da norma acima"
        )
    else:
        conclusao = (
            f"CONCLUSÃO: Foram evidenciadas indicações relevantes, "
            f"estando o ensaio {laudo_ext} segundo o critério de aceitação da norma acima"
        )

    mapping["{{CONCLUSAO}}"] = conclusao

    # Aplica substituições (corpo, cabeçalho, rodapé, shapes)
    replace_placeholders(doc, mapping)

    # ---------------- FOTOS ----------------
    foto1 = dados.get("FOTO_1")
    if foto1:
        inserir_foto_por_placeholder(doc, "{{FOTO_1}}", foto1, height_cm=10.0)

    foto2 = dados.get("FOTO_2")
    if foto2:
        inserir_foto_por_placeholder(doc, "{{FOTO_2}}", foto2, height_cm=4.0)

    # ---------------- SALVAR ----------------
    num = dados.get("NUMRELATORIO", "").strip() or "0000"
    filename = f"Relatório {num}-PM.docx"
    out_path = os.path.join(output_dir, filename)

    doc.save(out_path)
    return os.path.abspath(out_path)