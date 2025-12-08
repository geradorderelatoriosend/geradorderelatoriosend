import os
from datetime import date

from docx import Document

from utils_docx import (
    replace_placeholders,
    inserir_foto_por_placeholder,
    format_date_br,
    date_por_extenso,
)


def generate_lp_report(dados: dict, template_path: str, output_dir: str) -> str:
    """
    Gera o relatório de Líquido Penetrante (LP) a partir do TEMPLATE de LP.

    - dados: dicionário com os campos (NUMRELATORIO, EMPRESA, DATA_INSP, etc.)
    - template_path: caminho absoluto do LP_TEMPLATE.docx
    - output_dir: pasta onde o relatório será salvo
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template não encontrado: {template_path}")

    os.makedirs(output_dir, exist_ok=True)

    doc = Document(template_path)

    # ----------------------------------------------------
    # Mapeamento de placeholders de TEXTO (sem as FOTOS)
    # ----------------------------------------------------
    mapping: dict[str, str] = {}

    for key, value in dados.items():
        # NÃO mapear as fotos aqui, elas são tratadas separadamente
        if key in ("FOTO_1", "FOTO_2", "FOTO_3"):
            continue

        ph = f"{{{{{key}}}}}"  # ex: "DATA_INSP" -> "{{DATA_INSP}}"

        if key == "DATA_INSP" and isinstance(value, date):
            # Data no formato dd/mm/aaaa
            mapping[ph] = format_date_br(value)
        else:
            mapping[ph] = "" if value is None else str(value)

    # Data por extenso
    data_insp = dados.get("DATA_INSP")
    if isinstance(data_insp, date):
        mapping["{{DATA_INSP_EXTENSO}}"] = date_por_extenso(data_insp)
    else:
        mapping["{{DATA_INSP_EXTENSO}}"] = ""

    # ----------------------------------------------------
    # Conclusão automática (APROVADO / REPROVADO)
    # ----------------------------------------------------
    laudo = dados.get("LAUDO")
    laudo_ext = dados.get("LAUDO_EXTENSO") or ""

    if laudo == "A":
        conclusao = (
            f"CONCLUSÃO: Não Foram evidenciadas indicações relevantes, "
            f"estando o ensaio {laudo_ext} segundo o critério de aceitação da norma acima"
        )
    else:
        conclusao = (
            f"CONCLUSÃO: Foram evidenciadas indicações relevantes, "
            f"estando o ensaio {laudo_ext} segundo o critério de aceitação da norma acima"
        )

    mapping["{{CONCLUSAO}}"] = conclusao

    # Aplica substituições em todo o documento (corpo, cabeçalho, rodapé, shapes)
    replace_placeholders(doc, mapping)

    # ----------------------------------------------------
    # Inserção das fotos
    # ----------------------------------------------------
    foto1 = dados.get("FOTO_1")
    if foto1:
        inserir_foto_por_placeholder(doc, "{{FOTO_1}}", foto1, height_cm=10.0)

    foto2 = dados.get("FOTO_2")
    if foto2:
        inserir_foto_por_placeholder(doc, "{{FOTO_2}}", foto2, height_cm=4.0)

    foto3 = dados.get("FOTO_3")
    if foto3:
        inserir_foto_por_placeholder(doc, "{{FOTO_3}}", foto3, height_cm=4.0)

    # ----------------------------------------------------
    # Salvar arquivo
    # ----------------------------------------------------
    num = dados.get("NUMRELATORIO", "").strip() or "0000"
    filename = f"Relatório {num}-LP.docx"
    out_path = os.path.join(output_dir, filename)

    doc.save(out_path)
    return os.path.abspath(out_path)